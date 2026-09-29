"""Build editable PowerPoint, matching PDF, and English speaker notes.

Run with python-pptx, reportlab, matplotlib, numpy, and Pillow installed.
All scientific results come from the explicitly pinned local run below.
"""
from pathlib import Path
import csv, hashlib, json, math, re
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Ellipse, Circle
from matplotlib.mathtext import math_to_image
from matplotlib.font_manager import FontProperties
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from xml.sax.saxutils import escape
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
ASSETS=HERE/'assets'; ASSETS.mkdir(exist_ok=True)
RUN=ROOT/'result/appendix_d_20260921_235448_775259'
PAPER=ROOT/'overleaf_ntn_paper/direction.tex'
W,H=960,540
NAVY='#152C46'; TEAL='#087F8C'; BLUE='#326DD1'; GOLD='#D89B27'
INK='#263A50'; MUTED='#617487'; LIGHT='#EFF4F8'; PALE='#E5F3F2'; AMBER='#FFF3DD'; WHITE='#FFFFFF'; RED='#BC4B51'
for name,file in [('DV','DejaVuSans.ttf'),('DV-Bold','DejaVuSans-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/dejavu/'+file))
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':13,'axes.spines.top':False,
                     'axes.spines.right':False,'axes.labelcolor':INK,'text.color':INK,
                     'xtick.color':INK,'ytick.color':INK,'axes.edgecolor':'#C3CDD7',
                     'axes.titleweight':'bold','figure.facecolor':'white','savefig.facecolor':'white'})

def csvrows(relative):
    with (RUN/relative).open() as f:return list(csv.DictReader(f))

def saveplot(fig,name):
    fig.savefig(ASSETS/f'{name}.png',dpi=240,bbox_inches='tight')
    fig.savefig(ASSETS/f'{name}.pdf',bbox_inches='tight')
    plt.close(fig)
    return ASSETS/f'{name}.png'

# Replot actual CSV values for presentation readability; retain exact raw sources.
r=csvrows('E7/strategy_summary.csv')
fig,ax=plt.subplots(1,2,figsize=(10.8,3.8),gridspec_kw={'width_ratios':[1.5,1]})
styles={'F':(MUTED,'s','-'),'L':(GOLD,'x','--'),'T':(BLUE,'^','-'),'J':(TEAL,'o','-')}
for key in ['F','L','T','J']:
    rr=[a for a in r if a['strategy']==key]; color,marker,ls=styles[key]
    ax[0].plot([float(a['gamma_db']) for a in rr],[float(a['expected_bits_per_hz']) for a in rr],label=key,color=color,marker=marker,ls=ls,lw=2,ms=7)
ax[0].set(xlabel='INR limit Γ (dB)',ylabel='Expected cumulative service (bits/Hz)',xticks=[-15,-10,-5,0])
ax[0].legend(ncol=4,loc='upper left',fontsize=11); ax[0].grid(alpha=.2)
gammas=[-15,-10,-5,0];gains=[]
for g in gammas:
    v={a['strategy']:float(a['expected_bits_per_hz']) for a in r if float(a['gamma_db'])==g}
    gains.append(100*(v['J']/v['F']-1))
ax[1].bar([str(g) for g in gammas],gains,color=TEAL,width=.58)
for i,g in enumerate(gains):ax[1].text(i,g+.08,f'{g:.2f}%',ha='center',fontsize=12)
ax[1].set(xlabel='INR limit Γ (dB)',ylabel='J gain over F (%)',ylim=(0,3));ax[1].grid(axis='y',alpha=.2)
fig.tight_layout(w_pad=2.5);saveplot(fig,'finite_service')
rr=csvrows('fresh_rt_spatial/gamma_summary.csv');rr=[a for a in rr if a['gamma_db']]
fig,ax=plt.subplots(1,2,figsize=(10.8,3.6))
x=np.arange(4)
feas=[100*float(a['tn_static_feasible_fraction']) for a in rr]
powers=[100*float(a['median_power_fraction']) for a in rr]
ax[0].bar(x,feas,color=TEAL,width=.6)
for i,v in enumerate(feas):ax[0].text(i,v+3,f'{v:.1f}%',ha='center',fontsize=12)
ax[0].set(ylim=(0,120),ylabel='Sectors meeting TN SNR ≥ −6 dB (%)')
ax[1].semilogy(x,powers,'o-',color=BLUE,lw=2,ms=7)
for i,v in enumerate(powers):ax[1].annotate(f'{v:.3f}%',(i,v),xytext=(0,10),textcoords='offset points',ha='center',fontsize=12)
ax[1].set(ylabel='Median transmit power (% of Pmax)',ylim=(.01,3))
for a in ax:a.set(xticks=x,xticklabels=['−15','−10','−5','0'],xlabel='INR limit Γ (dB)');a.grid(axis='y',alpha=.2)
fig.tight_layout(w_pad=3);saveplot(fig,'rt_tradeoff')
a=json.loads((RUN/'E8/random_arrival/summary.json').read_text())
fig,ax=plt.subplots(figsize=(9.6,3.5));x=np.arange(3)
for j,(kind,color,label) in enumerate([('ul_active',TEAL,'UL-active arrival'),('ul_silent',GOLD,'UL-silent arrival')]):
    vals=[100*c['all_outage'] for c in a['cases'] if c['receiver_kind']==kind]
    xx=x+(j-.5)*.32;ax.bar(xx,vals,.3,color=color,label=label)
    for z,v in zip(xx,vals):ax.text(z,v+1.3,f'{v:.1f}%',ha='center',fontsize=12)
ax.set(xticks=x,xticklabels=['Prior envelope','No response','Delayed discovery'],ylabel='Active-time INR exceedance (%)',ylim=(0,65))
ax.legend(ncol=2,loc='upper center',fontsize=12);ax.grid(axis='y',alpha=.2);fig.tight_layout();saveplot(fig,'arrival')
a=json.loads((RUN/'E8/repeated_observation/summary.json').read_text())
fig,ax=plt.subplots(1,2,figsize=(10.3,3.4))
for i,c in enumerate(a['cases']):
    cor=c['accepted_indicator_correlation'];out=100*c['ntn0_outage']
    ax[0].bar(i,cor,color=[TEAL,BLUE][i]);ax[0].errorbar(i,cor,yerr=[[max(0,cor-c['accepted_correlation_ci_low'])],[max(0,c['accepted_correlation_ci_high']-cor)]],color=INK,capsize=5)
    ax[0].text(i,max(cor,0)+.1,f'{cor:.3f}',ha='center')
    ax[1].bar(i,out,color=[TEAL,BLUE][i]);ax[1].text(i,out+.35,f'{out:.2f}%',ha='center')
ax[0].set(ylabel='Acceptance correlation',ylim=(-.2,1.3));ax[1].set(ylabel='Foreground exceedance (%)',ylim=(0,10))
for z in ax:z.set(xticks=[0,1],xticklabels=['Independent','Correlated']);z.grid(axis='y',alpha=.2)
fig.tight_layout(w_pad=3);saveplot(fig,'correlation')
fig,ax=plt.subplots(figsize=(6.4,4.3));ax.set_aspect('equal');ax.axis('off')
ax.add_patch(Ellipse((0,0),4.8,1.5,angle=24,facecolor=PALE,edgecolor=TEAL,lw=2))
ax.add_patch(Ellipse((0,0),3.6,.45,angle=24,facecolor='#B7DAD9',edgecolor=TEAL,lw=1.5))
ax.plot([-2.8,2.8],[-1.25,1.25],color=TEAL,ls='--',lw=1)
ax.annotate('estimated path span',xy=(1.1,.5),xytext=(-2.6,1.9),arrowprops={'arrowstyle':'->','color':TEAL},fontsize=14)
ax.add_patch(Circle((1.35,.6),.49,fill=False,edgecolor=GOLD,lw=2))
ax.annotate('residual radius ρ',xy=(1.65,1),xytext=(.4,1.8),arrowprops={'arrowstyle':'->','color':GOLD},fontsize=14)
ax.scatter([1.45,-.7],[.95,-.4],c=[BLUE,TEAL],s=60,zorder=5)
ax.annotate('covered channel',xy=(1.45,.95),xytext=(.5,-1.4),arrowprops={'arrowstyle':'->','color':BLUE},fontsize=14)
ax.set(xlim=(-3.2,3.3),ylim=(-1.8,2.3));saveplot(fig,'uncertainty_set')
fig,ax=plt.subplots(figsize=(6,3.7));t=np.arange(13)
for color,vals,label in [(MUTED,.08+.028*t,'No refresh'),(TEAL,np.where(t<7,.08+.028*t,.08+.028*(t-1)),'Accepted at tick 7; reference = 1')]:
    ax.plot(t,vals,color=color,lw=2.5,label=label)
ax.axvspan(0,4,color=GOLD,alpha=.16);ax.axvspan(4,7,color=BLUE,alpha=.08)
ax.set(xlabel='Time (illustrative units)',ylabel='Uncertainty radius (schematic)',xticks=[0,1,4,7,12]);ax.legend(fontsize=9,loc='upper left');ax.grid(alpha=.2)
fig.tight_layout();saveplot(fig,'aging')

slides=[]
class Slide:
    def __init__(self,title,section='MODEL',source='Appendix D',take=None,dark=False):
        self.title,self.section,self.source,self.dark=title,section,source,dark
        self.e=[];self.notes='';slides.append(self)
        self.rect(0,0,W,H,NAVY if dark else WHITE)
        if not dark:
            self.text(section.upper(),44,24,850,13,TEAL,bold=True)
            title_size = min(30, 0.96 * 878 / max(1, pdfmetrics.stringWidth(title, 'DV-Bold', 30)) * 30)
            self.text(title,44,53,878,title_size,NAVY,bold=True)
            self.rect(44,104,55,4,TEAL)
        if take:self.banner(take)
    def rect(self,x,y,w,h,color=LIGHT,r=0,line=None):self.e.append(dict(kind='rect',x=x,y=y,w=w,h=h,color=color,r=r,line=line))
    def text(self,t,x,y,w=860,size=20,color=INK,bold=False,align='left',leading=1.22):
        self.e.append(dict(kind='text',t=t,x=x,y=y,w=w,size=size,color=color,bold=bold,align=align,leading=leading))
    def line(self,x1,y1,x2,y2,color=TEAL,width=2,dash=False):self.e.append(dict(kind='line',x=x1,y=y1,x2=x2,y2=y2,color=color,width=width,dash=dash))
    def arrow(self,x1,y1,x2,y2,color=TEAL,width=2):
        self.line(x1,y1,x2,y2,color,width)
        ang=math.atan2(y2-y1,x2-x1);a=8
        self.e.append(dict(kind='triangle',points=[(x2,y2),(x2-a*math.cos(ang-.5),y2-a*math.sin(ang-.5)),(x2-a*math.cos(ang+.5),y2-a*math.sin(ang+.5))],color=color))
    def image(self,path,x,y,w,h):self.e.append(dict(kind='image',path=str(path),x=x,y=y,w=w,h=h))
    def eq(self,formula,x,y,w=860,h=42,size=27,color=INK):
        key=hashlib.sha256((formula+str(size)+color).encode()).hexdigest()[:15]
        path=ASSETS/f'eq_{key}.png'
        with plt.rc_context({'savefig.transparent': True}):
            math_to_image('$'+formula+'$',str(path),prop=FontProperties(size=size),dpi=300,color=color)
        self.image(path,x,y,w,h)
    def banner(self,t,color=PALE):
        self.rect(44,454,872,46,color,r=8);self.text(t,60,465,842,17,NAVY,bold=True)
    def card(self,x,y,w,h,title,body,color=LIGHT):
        self.rect(x,y,w,h,color,r=10)
        head_size = min(21, 0.96 * (w-36) / max(1, pdfmetrics.stringWidth(title, 'DV-Bold', 21)) * 21)
        self.text(title,x+18,y+17,w-36,head_size,TEAL,bold=True)
        self.text(body,x+18,y+54,w-36,18)
    def bullets(self,items,x,y,w,size=20,gap=67):
        for i,t in enumerate(items):self.text('•',x,y+i*gap,16,size,TEAL,bold=True);self.text(t,x+23,y+i*gap,w-23,size)
    def note(self,t):self.notes=t.strip();return self

s=Slide('Service-Constrained Joint Sensing and Robust Transmission',dark=True,source='Paper: Appendix D · direction.tex | 22 September 2026')
s.text('APPENDIX D',52,45,850,16,'#6ED2CC',bold=True)
s.text('When to listen.\nHow to transmit.\nHow to keep service.',52,101,850,43,WHITE,bold=True,leading=1.18)
s.text('Joint sensing, robust beamforming and causal scheduling\nfor terrestrial / non-terrestrial coexistence',55,291,830,23,'#D7E4ED')
for i,t in enumerate(['Information','Protection','Service']):s.rect(55+i*272,393,247,54,'#24425C',r=8);s.text(t,70+i*272,410,217,20,WHITE,bold=True,align='center')
s.text('Theory + reproducible finite-model and fresh ray-tracing evidence',55,470,850,15,'#BCD0DD')
s.note('Opening: The main paper tells us where to suppress interference. Appendix D asks the next question: when should a base station pay the cost of sensing, and how should it transmit until the next useful observation? The aim is to maximize delivered terrestrial service while controlling deadline failures and non-terrestrial interference exposure. This presentation follows the five parts of the appendix, then adds explicitly labeled results from our current implementation. The theory is conditional on its observation and protection models. Current fresh ray tracing validates the spatial component, not the complete physical dynamic controller. Suggested duration: 25–30 minutes for slides 1–25; slides 26–30 are backups.')

s=Slide('A better estimate is useful only if it earns back its cost','MOTIVATION',take='The objective is delivered service after sensing overhead—not estimation accuracy alone.')
s.card(44,142,270,258,'Listen longer','Potentially better information\n\nMore lost DL/UL resources\n\nThe result arrives later',PALE)
s.card(345,142,270,258,'Transmit now','More immediate service\n\nOlder channel information\n\nPossible power backoff')
s.card(646,142,270,258,'Stay silent','Zero controlled-BS leakage\n\nNo TN downlink service\n\nTN uplink may continue',AMBER)
s.note('Explain the three competing actions. Listening is not free: reusing the front end removes scheduled downlink and uplink opportunities. Continuing service uses old, potentially less informative protection sets. Silence controls this base station’s interference contribution but sacrifices downlink bits. A received estimate is also already aged by the time processing finishes. The desired optimization compares the complete service and risk consequences of all three choices. This is why a static INR curve, or a smaller direction-estimation error, is not sufficient evidence for the joint controller.')

s=Slide('One controlled sector, two networks, partial information','SYSTEM',source='Appendix D.A–B · eq:joint_exposure')
s.card(350,183,250,110,'TN base station','Shared array and front end',PALE)
s.card(46,151,235,118,'TN user','Desired DL service\nFixed DL/UL calendar')
s.card(670,138,244,125,'NTN receiver','Protected DL reception\nMay emit observable UL',AMBER)
s.card(670,320,244,100,'UL-silent NTN','DL-active, no UL signal',AMBER)
s.arrow(350,223,281,210,BLUE);s.text('service',291,176,70,13,BLUE)
s.arrow(600,213,670,196,RED);s.text('interference',587,157,92,12,RED)
s.arrow(672,246,600,258,TEAL);s.text('UL sensing',581,273,99,12,TEAL)
s.arrow(600,270,670,355,RED)
s.text('Known online: current TN CSI, released observations, actions and service history.\nHidden: true NTN DL channels, identities, future activity and future observations.',46,355,570,17)
s.note('TN means terrestrial network; NTN means non-terrestrial network. The controlled object is one TN sector with one scheduled stream per resource. Its downlink can interfere with an NTN downlink receiver. The base station listens to an NTN uplink band, which can differ from the protected downlink frequency. An uplink-silent terminal can still be receiving its satellite downlink and therefore require protection. The controller has its own current TN channel in the declared model, but it does not receive the true victim channel or the future physical state. Other interference is fixed or separately budgeted in the exact core.')

s=Slide('The appendix connects sensing to service through two optimizers','ARCHITECTURE',take='SOCP chooses the best protected transmission; LP decides when it is worth using.')
for i,(title,body) in enumerate([('Available history','Released estimates\nQuality and age'),('Protection catalog','Anonymous path groups\nResidual and gain bounds'),('Robust SOCP','Beam + power\nBest feasible service')]):
    s.card(44+i*299,140,274,130,title,body,PALE if i==2 else LIGHT)
    if i<2:s.arrow(321+i*299,206,338+i*299,206)
s.card(191,320,579,101,'Finite-history occupation-measure LP','Choose listening time, duration, template, service and silence',PALE)
s.arrow(779,276,712,316);s.arrow(204,317,174,276)
s.text('observation transitions + risk + deadline costs',273,287,480,15,MUTED,align='center')
s.note('Walk left to right. Released sensing information is mapped to a catalog of uncertainty sets and conditional coverage-risk bounds. For each usable state, a robust second-order cone program computes the best desired-channel amplitude and the corresponding power. The outer linear program uses those service values, sensing outcomes and deadline events to choose a randomized causal policy. Observation quality, delay and unsuccessful sensing enter the transition model. These are layers of one design, not three unrelated algorithms. This slide is the roadmap: timing, sets, service and risk, SOCP, then LP.')

s=Slide('A new estimate is unavailable until its release time','A · TIMING',source='Appendix D.A · eq:joint_timing · illustrative 1/2/1/3 time units',take='At release: age = 7 − 1 = 6. A rejected observation keeps the previous record.')
x0=110;scale=91
for start,length,label,color in [(0,1,'Enter',GOLD),(1,2,'Listen',TEAL),(3,1,'Return',GOLD),(4,3,'Process',BLUE)]:
    s.rect(x0+start*scale,166,length*scale-3,63,color,r=4);s.text(label,x0+start*scale+3,186,length*scale-9,17,WHITE,bold=True,align='center')
for t in [0,1,3,4,7]:s.line(x0+t*scale,232,x0+t*scale,242,MUTED,1);s.text(str(t),x0+t*scale-17,246,34,15,MUTED,align='center')
s.text('RF unavailable: no affected TN DL/UL',111,287,350,17,RED,bold=True)
s.text('RF returned: service may resume\nusing the old record',477,287,329,17,BLUE,bold=True)
s.eq(r't_{\rm ref}=s+g^{\rm in},\qquad t_{\rm use}=s+g^{\rm in}+\tau+g^{\rm out}+t_p',58,349,842,34,25)
s.eq(r'\Delta_0=\tau+g^{\rm out}+t_p',238,401,490,31,25)
s.note('Use the example rather than starting with symbols. Entry takes one time unit, collection two, return one, and processing three. Samples are referenced to time one, not time seven. The front end is unavailable until time four. With nonblocking processing, the base station can resume calendar-permitted service at time four, but it must use the old record until time seven. Acceptance at time seven gives a record aged six units. Rejection preserves the old reference time. This is an illustrative appendix timeline, not the shorter guard values in the current 12-tick simulator. Blocking processing would extend unavailability through release.')

s=Slide('Longer listening trades accuracy against usable service time','A–B · INFORMATION AGE',source='Appendix D.A–B · schematic, not fitted experimental data',take='The policy must compare successful, rejected and no-detection branches.')
s.image(ASSETS/'aging.png',37,133,540,292)
s.bullets(['Age grows from the sample reference time.','An accepted refresh may shrink uncertainty.','Longer sensing also delays release and consumes service time.'],594,153,320,20,gap=88)
s.note('This plot is deliberately schematic. It visualizes uncertainty growing while the previous record is used; after an accepted observation, the reference changes to the start of that listening window, not to its release time. A refreshed record can therefore start with substantial age. Better observation quality can reduce an estimation component, while age, calibration error and missing multipath remain. The appendix does not assume a universal inverse-square-root law in listening duration. For a fixed ground BS-to-VSAT link, satellite orbital speed must not be substituted for ground-link direction change.')

s=Slide('Protect a set of complete channels, not individual peaks','B · UNCERTAINTY',source='Appendix D.B · eq:joint_multipath_set',take='Complex coefficients retain coherent path addition; the full receiver channel must be covered.')
s.eq(r'\mathcal{G}_r=\{\mathbf{A}_r\mathbf{c}+\mathbf{e}:\ \|\mathbf{c}\|_2\leq C_r,\ \|\mathbf{e}\|_2\leq\rho_r\}',48,135,868,46,29)
s.image(ASSETS/'uncertainty_set.png',40,192,490,239)
s.text('Aᵣ: reconstructed DL path responses\n\nCᵣ: bound on complex coefficients\n\nρᵣ: unresolved paths and model error\n\nr: anonymous group, not a UE identity',549,210,352,19)
s.note('Explain A, C and rho separately. Columns of A are downlink array responses reconstructed from observed directions. The complex coefficient vector permits arbitrary coherent relative phases within its norm bound. The residual ball covers unmodeled or nontransferable paths and other errors. A group does not require identification of a specific receiver. A complete victim channel must belong to at least one imposed group; independently protecting individual peaks is insufficient when their coherent sum is the actual channel. The two-dimensional figure is a conceptual projection of a complex high-dimensional set, not an exact physical channel plot.')

s=Slide('The worst-case leakage has an exact closed form','B · ROBUST PROTECTION',source='Appendix D.B · Proposition: exact set-wise protection',take='Exact for the declared set; the set itself may conservatively overbound physical channels.')
s.eq(r'\sup_{\mathbf{f}\in\mathcal{G}_r}|\mathbf{f}^{H}\mathbf{v}|=C_r\|\mathbf{A}_r^{H}\mathbf{v}\|_2+\rho_r\|\mathbf{v}\|_2',48,150,864,57,31)
s.card(44,254,422,129,'Modeled-path contribution','Suppress the estimated DL subspace.\nMore directions can cost spatial freedom.',PALE)
s.card(494,254,422,129,'Residual contribution','Uncertain directions remain.\nPower backoff limits their leakage.',AMBER)
s.eq(r'C_r\|\mathbf{A}_r^{H}\mathbf{v}\|_2+\rho_r\|\mathbf{v}\|_2\leq\sqrt{\Gamma}',160,404,640,32,25)
s.note('The upper bound comes from Cauchy–Schwarz and the triangle inequality. It is attainable: choose the coefficient vector aligned with A Hermitian v and the residual aligned with v, with phases that add rather than cancel. That is why the support expression is exact for this uncertainty set. The two terms explain the engineering tradeoff. Steering away from the estimated subspace reduces the first term, but a residual ball generally leaves a term proportional to total beam norm. The INR threshold Gamma is a power ratio; therefore the amplitude constraint uses square root Gamma.')

s=Slide('A silent receiver can still force power backoff','B–C · BACKGROUND PROTECTION',source='Appendix D.B–C · eq:joint_exposure · background catalog',take='Passive UL sensing cannot distinguish an absent receiver from a permanently silent one.')
s.eq(r'\mathbf{v}=\sqrt{p}\,\mathbf{w},\quad\|\mathbf{w}\|_2=1,\quad\mathrm{INR}_i=|\mathbf{f}_i^H\mathbf{v}|^2',57,140,846,45,28)
s.eq(r'\mathbf{f}_i=\sqrt{\chi_i P^{\max}/N_i}\,\mathbf{g}_i^{\rm D}',247,190,461,28,24)
s.card(44,223,413,189,'Directional information','A known path span supports spatial suppression.\n\nUnknown complex DL gains still require valid bounds.',PALE)
s.card(488,223,428,189,'Arbitrary-direction background','If ||f|| ≤ β, every direction is possible.\n\nThe only universal control is beam norm.',AMBER)
s.eq(r'\beta\|\mathbf{v}\|_2\leq\sqrt{\Gamma}\quad\Rightarrow\quad p\leq\min\{1,\Gamma/\beta^2\}',503,374,394,29,22)
s.note('The normalized channel f includes receiver noise, spectral coupling and the power limit. The vector v combines beam shape and power fraction: physical power is Pmax times p. A background norm ball can be represented either with A equal to the identity and coefficient bound beta, or as a pure residual ball of radius beta; both give beta times beam norm. A receiver with no observable uplink requires prior coverage, a valid arrival/activity model, or side information. The method cannot infer an unknown downlink antenna gain from a MUSIC weight alone. Zero downlink transmission protects the controlled contribution, but not unrelated external interference.')

s=Slide('Service reliability and interference exposure use different metrics','C · PERFORMANCE METRICS',source='Appendix D.C · eq:joint_service_target · eq:joint_activity_outage',take='Use identical windows, calendars, references and budgets for every policy.')
s.card(44,135,423,275,'TN: deadline service','Accumulate only available resources.\nCheck each user, direction and window.\nThe demand is frozen in advance.',PALE)
s.eq(r'\Pr\{B_{u,x,m}<L_{u,x,m}\}\leq\delta^{\rm TN}_{u,x,m}',65,304,380,47,25)
s.text('Mean throughput cannot replace this event.',64,373,380,16,TEAL,bold=True)
s.card(494,135,422,275,'NTN: active-time exposure','Include missed and UL-silent receivers.\nCount activity during gaps and silence.\nAggregate durations across episodes.',LIGHT)
s.eq(r'\mathcal{O}_{i,m}=\frac{\mathbb{E}[\int_{W_m}z_i(t)\,\mathbf{1}\{\mathrm{INR}_i>\Gamma\}\,dt]}{\mathbb{E}[\int_{W_m}z_i(t)\,dt]}',510,298,390,62,25)
s.text('A ratio of expectations, not a window event.',511,373,387,16,BLUE,bold=True)
s.note('The TN constraint is a chance constraint on deadline service, independently for each applicable user, direction and window. The reference and demand cannot be recalculated after seeing a favorable Gamma or future channel. The NTN metric is the expected exceeding active duration divided by expected active duration. In measurements, estimate it by total exceeding duration divided by total active duration across complete independent episodes, not an unweighted average of episode ratios with varying denominators. BS silence makes the controlled interference zero but does not erase the receiver’s active time. Gamma controls interference magnitude; delta NTN controls the allowable active-time fraction.')

s=Slide('Coverage risk bridges robust constraints and actual outage','C · CONDITIONAL RISK',source='Appendix D.C/E · eq:joint_risk_union · eq:joint_risk_cost',take='Marginal calibration coverage is not a certificate for every adaptively selected history.')
s.eq(r'\Pr\{E_i(t)\mid\mathcal{H}_k,a,z_i(t)=1\}\leq\eta_i',53,142,856,46,29)
for i,(head,body) in enumerate([('Coverage failure','Missing receivers\nor missing groups'),('Set failure','Channel outside the\nmodeled uncertainty set'),('Gain-bound failure','Noise / gain / coupling\nbounds do not hold')]):s.card(44+i*299,218,274,121,head,body)
s.eq(r'\eta_i=\min\{1,\eta_i^{\rm cov}+\eta_i^{\rm set}+\eta_i^{\rm gain}\},\qquad c_{i,m}=\eta_i d_{i,m}',67,370,826,44,27)
s.note('On valid catalog coverage, the robust beam constraint makes the exceedance event impossible for the controlled contribution. Exceedance can therefore occur only through a coverage or bound failure. Use component upper bounds under the same history, action and receiver-activity conditioning; their union bound does not require independence. Multiplying a valid throughout-segment conditional probability bound by expected active duration yields a risk-duration coefficient. Overall average calibration coverage does not automatically remain valid after a policy selects particular histories. Sparse histories require conservative fallback bounds or a narrower model. Statistical confidence in an estimated bound is separate from the physical outage budget.')

s=Slide('The joint problem maximizes net service under explicit budgets','C · JOINT PROBLEM',source='Appendix D.C · eq:joint_policy_problem',take='The old leakage penalty λ is replaced by physical constraints and service requirements.')
s.rect(44,135,872,279,LIGHT,r=12)
s.eq(r'\max_{\Pi\ {\rm causal}}\quad\mathbb{E}_{\Pi}\!\left[\sum_{u,x,m}B_{u,x,m}\right]',97,154,760,52,31)
s.eq(r'\Pr_{\Pi}\{B_{u,x,m}<L_{u,x,m}\}\leq\delta^{\rm TN}_{u,x,m}',105,228,735,41,28)
s.eq(r'C_{i,m}(\Pi)\leq\delta_i^{\rm NTN}D_{i,m}(\Pi)',160,289,628,39,28)
s.text('Calendar • RF availability • transmit power • robust set constraints',97,355,763,20,NAVY,align='center')
s.note('The decision variable is a causal policy, not just a beam vector. It chooses legal sensing starts, listening durations and templates, plus transmission or silence. C is a computable upper bound on expected exceeding active duration; D is expected active duration. For valid positive activity, C less than delta times D is sufficient for the actual outage target. The objective directly rewards bits delivered after overhead, so a common arbitrary penalty lambda is unnecessary in this core. Infeasibility is a legitimate result: do not relax the TN requirement silently just to draw a curve.')

s=Slide('At one state, a convex program selects beam and power','D · ROBUST SOCP',source='Appendix D.D · eq:joint_value_socp',take='This solves the best current protected transmission—not the sensing schedule.')
s.rect(44,135,528,280,PALE,r=12)
s.eq(r'V_k=\max_{\mathbf{v}}\ \mathrm{Re}\{\mathbf{h}^{H}\mathbf{v}\}',63,164,487,47,29)
s.eq(r'\mathrm{Im}\{\mathbf{h}^{H}\mathbf{v}\}=0,\quad\|\mathbf{v}\|_2\leq1',63,233,487,36,25)
s.eq(r'C_r\|\mathbf{A}_r^H\mathbf{v}\|_2+\rho_r\|\mathbf{v}\|_2\leq\sqrt{\Gamma}',63,294,487,42,25)
s.text('for every imposed protection group r',76,365,459,17,MUTED,align='center')
s.bullets(['Rotate the common phase: no loss of optimality.','Recover p = ||v||² and the beam direction.','Convert Vₖ into service; retain mute separately.'],603,149,306,20,gap=89)
s.note('With fixed current TN channel and protection catalog, all constraints depend on phase-invariant magnitudes. Rotate the beam so the desired inner product is real and nonnegative, then maximize its real part. This gives an SOCP. The optimum can use full power, partial nulling, or power backoff. When the solution is nonzero, p is its squared norm and w is the unit-normalized direction. Its rate is B log2(1 + Pmax V squared divided by noise plus external interference). The feasible beam set always contains zero; positive service feasibility is the real issue. Age expansion of fixed-center uncertainty sets can only shrink the feasible set.')

s=Slide('Why the outer controller does not search every beam','D · BEAM ELIMINATION',source='Appendix D.D · Proposition: lossless beam elimination',take='Keep one SOCP maximizer and an explicit mute action at each available history.')
s.card(44,145,420,236,'Coupling argument','Keep the same physical trajectory.\nKeep sensing and mute decisions.\nReplace each nonmute beam by the SOCP optimizer.\n\nDelivered bits cannot decrease.',PALE)
s.card(494,145,422,236,'Conditions that make it valid','Fixed calendar and external interference.\nNo energy carryover or switching cost.\nExogenous physical/observation laws.\nOne shared certified nonmute risk bound.',LIGHT)
s.text('The result concerns certified risk; beam-dependent true-outage costs need a different reduction.',52,404,853,17,RED,bold=True)
s.note('This is the key reduction linking continuous beam design to a finite policy problem. Consider a virtual copy of any original policy with its original service counters. On the same exogenous trajectory and random seed, keep the sensing and mute choices, but use the SOCP maximizing beam whenever transmitting. Actual service only increases, while the certified risks and future observation laws remain unchanged under the stated assumptions. Minimum-service failures cannot become worse. A virtual copy is needed because later decisions may depend on the original service counters. This theorem does not cover beam-dependent actual outage costs, intertemporal energy, coupled network SINR, or uncertain current TN CSI without further analysis.')

s=Slide('A causal state contains what is known—not the hidden channel','E · HISTORY GRAPH',source='Appendix D.E · eq:joint_history_state · eq:joint_belief_update',take='Pending metadata includes the release time, but never the unreleased estimate.')
s.text('S = time + revealed history + RF / pending status + delivered bits + record timestamps',48,133,864,19,NAVY,bold=True)
s.card(52,221,210,112,'Known history','Legal actions\nAvailable record',PALE)
s.card(378,163,216,95,'Serve / mute','Use available information')
s.card(378,321,216,95,'Listen','Commit resources')
s.arrow(265,250,373,208);s.arrow(265,280,373,355)
for y,lab,col in [(280,'Accepted → new record',TEAL),(344,'Rejected → old record',GOLD),(408,'No detection → old record',MUTED)]:
    s.rect(691,y-16,225,43,PALE if col==TEAL else LIGHT,r=6);s.text(lab,700,y-5,207,15,col,bold=True);s.arrow(598,367,687,y+4,col)
s.text('At release only',686,228,224,17,TEAL,bold=True)
s.note('The complete observable history supports a belief over hidden physical states without revealing them to the controller. The state also carries timing, pending-release metadata, delivered service and timestamps, because those affect legal actions and deadline events. The graph branches on observations actually released. Rejection and no detection remain explicit branches; pruning them and renormalizing changes the declared model. Time increases along each edge, so the graph is acyclic. It can nevertheless grow exponentially with horizon and observation alphabet, which is why the exact solver is a small-instance benchmark rather than an automatic solution to arbitrary large deployments.')

s=Slide('Occupation measures turn policy optimization into an LP','E · LINEAR PROGRAM',source='Appendix D.E · eq:joint_occupancy_objective through eq:joint_policy_extraction',take='Normalize state–action mass to recover the executable randomized policy.')
s.text('x(S,a) = probability of reaching state S and choosing action a',49,132,866,20,TEAL,bold=True)
s.eq(r'\max_{x\geq0}\ \sum_{S,a}x(S,a)\,r(S,a)',70,179,817,42,29)
s.eq(r'\sum_a x(S,a)=\mu_0(S)+\sum_{\bar S,\bar a}P(S\mid\bar S,\bar a)x(\bar S,\bar a)',70,241,817,46,26)
s.eq(r'\sum_{S,a}x(S,a)\bar g_{u,x,m}(S,a)\leq\delta^{\rm TN}_{u,x,m}',72,305,813,39,26)
s.eq(r'\sum_{S,a}x(S,a)\,[c_{i,m}(S,a)-\delta_i^{\rm NTN}d_{i,m}(S,a)]\leq0',72,362,813,43,26)
s.eq(r"\pi^\star(a\mid S)=x^\star(S,a)/\sum_{a'}x^\star(S,a')",254,413,451,27,22)
s.note('Read the four lines as reward, probability flow, TN deadline failures and NTN risk-duration budgets. The LP variables are probabilities of state-action visits. Flow conservation ensures they correspond to a causal policy. The coefficient g-bar is a deadline-failure event cost, charged once after crediting the just-completed resource. Coefficients c and d are expected exceeding-duration bounds and activity durations under the declared model. All SOCP values and coefficients are fixed before solving the LP. At positive-mass states, divide each action mass by the sum over actions to get action probabilities. Terminal states absorb mass and have no outgoing actions.')

s=Slide('Randomization is sometimes required, not merely convenient','E · POLICY EXECUTION',source='Appendix D.E + E6/E6_validation.csv · illustrative mixture below',take='Executing only the most likely action can destroy the probability constraints.')
s.card(44,138,262,135,'Policy A (illustration)','NTN cost = 0.02\nTN failure = 0.20',AMBER)
s.card(353,138,262,135,'Policy B (illustration)','NTN cost = 0.08\nTN failure = 0.02',LIGHT)
s.card(663,138,253,135,'50 / 50 mixture','NTN cost = 0.05\nTN failure = 0.11',PALE)
s.text('Example limits: NTN cost ≤ 0.05 and TN failure ≤ 0.11.\nA fails TN, B fails NTN; the mixture satisfies both.',53,304,850,21)
s.rect(44,382,872,48,LIGHT,r=7)
s.text('Actual E6 check: 128 deterministic policies; none feasible in the randomization case; mixed optimum = 1.125.',58,395,844,15,NAVY,bold=True)
s.note('The top example uses deliberately simple illustrative numbers, not measured results. Treat the NTN cost as a normalized risk cost with a common fixed active denominator. Policy A is safe but has excessive TN failures; policy B has better service reliability but excessive NTN risk. The equal mixture is feasible in both coordinates. A deterministic choice of either whole policy fails a requirement. In the actual E6 finite test, all 128 deterministic policies are infeasible for the randomization-required configuration, but their randomized mixture and the occupation LP achieve the same feasible objective 1.125. The extracted probabilities must be executed as probabilities; taking an argmax is not an equivalent implementation.')

s=Slide('There are three distinct levels of guarantee','THEOREM & SCOPE',source='Appendix D.B/D/E · robust support, beam elimination, finite optimality theorem')
s.card(44,137,872,85,'1  Set-wise protection','Every channel inside an imposed set satisfies the INR limit.',PALE)
s.card(44,239,872,85,'2  Finite-model optimality','Exact SOCP + complete finite history + exact transitions/costs → optimal certified policy.',LIGHT)
s.card(44,341,872,93,'3  Physical-system protection','Also requires valid physical state laws, active-time expectations and conditional certificates.',AMBER)
s.note('Do not merge these claims. The support-function proposition is exact for the declared uncertainty set. The beam-elimination proposition and occupation-measure theorem establish optimality over randomized causal policies in the declared finite certified problem, assuming exact transitions, coefficients and solutions. A physical outage guarantee additionally requires that the implemented model and conditional certificates remain valid for the physical system. A favorable static CDF alone does not establish that final condition. Similarly, approximate history compression or numerical solutions require their approximation and residual errors to be addressed. State these distinctions explicitly during the talk.')

s=Slide('Most optimization is offline; execution follows released events','IMPLEMENTATION',source='Appendix D.E · finite-model construction and execution',take='The complete history can be exponential even though the explicit LP is polynomial in size.')
for i,(head,body) in enumerate([('Freeze the model','Calendar, demands,\nkernels and risk bounds'),('Build and solve','Reachable histories\nSOCP values → LP'),('Execute causally','Sample an action\nReveal only released results')]):
    s.card(44+i*299,158,274,172,head,body,PALE if i==2 else LIGHT)
    if i<2:s.arrow(320+i*299,242,339+i*299,242)
s.text('Carry pending jobs, record ages and service counters through deadlines.\nCharge online computation latency. Report infeasibility separately from solver failure.',59,366,838,20)
s.note('The algorithm freezes model inputs, expands all reachable histories, builds catalogs and risk bounds, solves or caches state-specific SOCPs, computes transition rewards and costs, and solves the LP. Online operation follows the conditional policy and only updates on events that have actually been released. A beam cache key must include every quantity affecting the solution, including current CSI, sets, radii, interference and Gamma. Deadline boundaries do not justify resetting ages, pending jobs or unfinished requirements. If a zero-probability history is physically reached, that signals model mismatch: a silent fallback controls the BS contribution, but its service loss must still be counted.')

s=Slide('What the current experiments validate','EVIDENCE MAP',source='Run 20260921_235448_775259 · E4–E8 + fresh_rt_spatial')
rows=[('E4','Set calibration and coverage','Synthetic whole-scene split'),('E5','Worst leakage and robust SOCP','Numerical identities + geometry tests'),('E6','Beam reduction and policy LP','Complete-policy enumeration'),('E7','F / T / L / J service tradeoff','Declared finite dynamic model'),('E8','Failure, delay, arrivals, correlation','Model stress + fixed causal diagnostics'),('Fresh RT','New drops + DL/UL tracing + Γ-CDF','Static held-out physical channels')]
y=137
for i,(a,b,c) in enumerate(rows):
    s.rect(44,y,872,46,PALE if i==5 else (LIGHT if i%2==0 else WHITE),r=4)
    s.text(a,58,y+13,108,17,TEAL,bold=True);s.text(b,174,y+13,417,16);s.text(c,592,y+13,310,14,MUTED);y+=50
s.note('This table prevents a common overclaim. E4 verifies a synthetic whole-scene calibration procedure; it does not supply a physical conditional risk table for arbitrary adaptive histories. E5 and E6 provide solver and reduction checks. E7 is the dynamic finite-model service comparison. E8 includes both optimized-model stress cases and separate fixed-policy mechanism diagnostics. Fresh ray tracing now really redraws user positions and recomputes propagation, but it validates the static spatial component. The 34 regression tests and full notebook execution support implementation consistency, not a universal physical deployment guarantee. The pinned run is September 21 at 23:54:48 local time.')

s=Slide('Finite-model gains are positive, but small','RESULTS · E7',source='E7/strategy_summary.csv · exact model expectations; Monte Carlo: 400 episodes per setting',take='J = T at three Γ values; current evidence does not show a broad duration-adaptation gain.')
s.text('F: fixed   T: timing only   L: duration only   J: joint adaptation',50,129,860,17,MUTED)
s.image(ASSETS/'finite_service.png',46,159,868,269)
s.note('All four classes use the same robust transmission subproblem and can serve or mute; fixed baselines are optimized within their declared classes. These are exact model expectations accumulated over 12 ticks, not per-tick throughput or noisy simulation means. At minus 15 dB, J yields 23.0089 bits/Hz versus F at 22.4713, a 2.39% gain. Its advantage over T is only 0.247%. At the other three Gamma values J and T coincide; F and L coincide throughout. Increasing the Monte Carlo sample count will tighten uncertainty but will not separate identical model optima. Sixteen main configurations were feasible. Confidence on actual simulated risk is discussed in a backup slide.')

s=Slide('Fresh ray tracing reveals a substantial service cost','RESULTS · PHYSICAL SPATIAL TEST',source='fresh_rt_spatial/gamma_summary.csv · 6 drops, 2/2/2 split; 24 test sectors',take='Low INR alone is insufficient: the current empirical envelope is conservative.')
s.text('7 GHz DL / 6.3 GHz UL · 64 antennas · depth 3 · 800-snapshot MDL-MUSIC',50,129,863,16,MUTED)
s.image(ASSETS/'rt_tradeoff.png',45,161,870,265)
s.note('Unlike the earlier cache-only path, this run creates six fresh TN/NTN drops and retraces both bands. Each drop contains 12 TN users, 20 NTN receivers and 12 sectors. Two complete drops each are used for training, calibration and test, with all 24 test sectors evaluated. The plotted INR is one controlled sector’s contribution, not the network sum. At Gamma minus 15 dB, only 58.33% of sectors meet the static minus 6 dB TN SNR floor; median power is just 0.032% of Pmax. The broad empirical residual bound strongly limits power. Zero observed exceedances in this small test do not certify a tail probability. This is physical static validation, not physical dynamic E7/E8. Quick RT also uses a coarser search and ray budget than full.')

s=Slide('Unknown arrivals expose the gap before usable information','RESULTS · CAUSAL ARRIVAL DIAGNOSTIC',source='E8/random_arrival/summary.csv · fixed policies, 400 paired episodes per case',take='UL-silent arrivals require prior coverage; delayed discovery cannot protect the unseen past.')
s.image(ASSETS/'arrival.png',49,136,863,245)
s.text('Only released detections reach the controller.\nUL-active case: detection 67.5%; observed exceedance after protection activates = 0.',56,380,849,16)
s.text('Fixed diagnostics violate the main TN QoS budget; these are not optimized J policies.',56,426,849,14,RED,bold=True)
s.note('Arrival times are drawn by the evaluator and hidden from the controller. All three comparison policies share arrival trajectories, sensing random numbers and RF costs. The prior-envelope policy protects the newcomer class from the outset. No-response ignores detection. Delayed-discovery starts protection only after detection has completed and its processing delay has elapsed. For an UL-active newcomer, exposure falls from 47.53% with no response to 24.45% with delayed discovery; detection probability is 67.5%. The permanently UL-silent newcomer is never detected and receives no discovery-based benefit. These are mechanism diagnostics using fixed policies, not an optimal random-arrival LP solution. Their TN deadline failures violate the main budget, so they cannot be presented as feasible service improvements.')

s=Slide('Repeated observations now genuinely test temporal correlation','RESULTS · CORRELATION DIAGNOSTIC',source='E8/repeated_observation/summary.json · fixed two-listen policy, 400 paired episodes',take='Correlation is present; this run does not show a statistically significant performance loss.')
s.image(ASSETS/'correlation.png',48,143,865,261)
s.text('Every path executes two listens. The fixed diagnostic violates the main TN budget;\nit tests the observation mechanism, not a feasible optimized-J service gain.',53,409,856,16,RED)
s.note('The original optimized J policy listened at most once along each reachable trajectory, so making latent observation uniforms correlated did not actually test repeated observations. The new diagnostic fixes two listening occasions before observing any outcomes. Independent and correlated cases use paired physical trajectories and the same observation marginal probability. The first/second acceptance correlation changes from about minus 0.029, with a confidence interval spanning zero, to one. Foreground exposure is about 7.46% versus 7.38%; the paired difference confidence interval spans zero. Do not force a negative result into a robustness claim: the mechanism is now exercised, but these parameters do not show a significant loss. The fixed schedule is explicitly not feasible under the original TN constraints.')

s=Slide('The contribution is a causal connection from sensing to service','TAKEAWAYS',source='Appendix D + current implementation scope')
s.card(44,139,872,82,'One channel-set model','Single paths, coherent multipath and background uncertainty share a robust constraint.',PALE)
s.card(44,236,872,82,'One structured optimization','SOCP supplies protected service values; an occupation LP chooses a causal policy.',LIGHT)
s.card(44,333,872,101,'One clear validation boundary','Finite-model optimality is established conditionally; physical dynamic calibration remains to be built.',AMBER)
s.note('Close with three points. First, complete-channel uncertainty sets unify dominant-path and coherent-multipath protection, including unobserved background. Second, under explicit assumptions, beam elimination lets a finite causal policy combine sensing decisions, power, service and silence. Third, exact finite-model optimality and physical-system protection are different claims. We have verified numerical identities and policy consistency, exercised the previously missing causal mechanisms, and connected fresh physical ray tracing for spatial evaluation. The next research step is an independently calibrated physical dynamic model with repeated observations, user activity and timing, together with less conservative but valid observable-condition protection sets.')

# Backup slides
s=Slide('How listening quality enters the multipath residual','BACKUP · B',source='Appendix D.B · eq:joint_error_envelope · eq:joint_multipath_radius',take='Raw UL MUSIC weights do not certify protected-band gain or noise bounds.')
s.eq(r'\mathbf{f}_i=(\mathbf{A}_r+\Delta\mathbf{A}_r)\mathbf{c}+\mathbf{h}_r^{\rm tail}',65,139,826,44,29)
s.eq(r'\rho_r=C_r\left(\sum_p\varepsilon_{rp}^{\,2}\right)^{1/2}+\rho_r^{\rm tail}',91,204,777,51,30)
s.eq(r'\varepsilon(\tau,\Delta,\upsilon)=L[r_0(\tau,\upsilon)+\nu\Delta]+e_{\rm cal}+e_{\rm mp}',57,282,847,44,27)
s.text('Finite observation error + directional aging + calibration / multipath mismatch.\nCᵣ requires path-power bounds or independent calibration: coherent cancellation\ncan hide large path coefficients in a small total channel.',59,354,844,19)
s.note('Derive the residual bound by writing the error as delta-A times c plus the missing tail. The spectral norm of delta-A is no larger than its Frobenius norm, which is bounded by the root-sum-square of per-path response errors. Multiplying by C and adding the tail gives a sufficient radius. This is a constructive sufficient bound, not a claim that every path can be resolved. Total channel power cannot in general bound path coefficient norm, because coherent cancellation can hide large coefficients. The illustrative age envelope can be replaced by independently calibrated tables, and its validity must cover the intended service interval.')

s=Slide('Silence is feasible; positive deadline service may not be','BACKUP · D',source='Appendix D.D · eq:joint_service_feasibility · eq:joint_power_cap',take='This is a constant-input screening bound, not a universal per-symbol QoS constraint.')
s.eq(r'V_k\geq\sqrt{\frac{N+I}{P^{\max}}\left(2^{d/(BT_{\rm rem})}-1\right)}',99,148,760,60,31)
s.text('d remaining bits; Trem available DL time; fixed channel and catalog.',69,224,820,19,MUTED,align='center')
s.eq(r'\ell_r=C_r\sqrt{\lambda_{\min}(\mathbf{A}_r\mathbf{A}_r^H)}+\rho_r',81,279,800,45,28)
s.eq(r'p_{\rm cap}=\min\{1,\min_{r:\ell_r>0}\Gamma/\ell_r^2\},\quad V_k\leq\|\mathbf{h}\|_2\sqrt{p_{\rm cap}}',56,350,849,49,27)
s.note('For a fixed channel and catalog over remaining usable downlink time, completing d bits is equivalent to the amplitude threshold on the first line. The second pair gives a necessary screening bound using the smallest eigenvalue of the protected subspace Gram matrix and the residual radius. A rank-deficient span can leave a spatial null space, but nonzero residual still caps power. A full span can cap power even when residual is zero. The screening bound is not sufficient, because it ignores alignment of the desired TN channel with the protected subspace. Do not use a current-catalog bound to discard a full sensing plan whose future accepted observation can change that catalog.')

s=Slide('A network extension needs local budgets and an external reserve','BACKUP · NETWORK INTERFACE',source='Appendix D.E · eq:joint_network_allocation · eq:joint_external_budget',take='Independent data streams add average powers; failure events need not be independent.')
for i in range(3):
    s.card(44+i*299,139,274,100,f'Sector {i+1}','Local amplitude/risk budgets',PALE)
s.eq(r'\sum_b\Gamma_b\leq\Gamma^{\rm net},\qquad\sum_b\delta_{ib}\leq\delta_i^{\rm net}',66,284,828,49,30)
s.eq(r'\Gamma^{\rm net}=\Gamma^{\rm total}-\overline{J}^{\rm ext}',164,356,635,44,29)
s.note('For independent BS data streams, symbol-averaged interference powers add. Assign local Gamma budgets whose sum is within the network budget and local risk allowances whose sum is within the per-receiver risk target. If every local contribution is below its budget, the total is below the summed budget, so network exceedance requires at least one local exceedance. A union bound works under common receiver-activity conditioning without independent certificate failures. Uncontrolled TN uplink or other sources require a valid reserved external contribution, including failure risk if the reserve is uncertain. Endogenous inter-sector TN SINR coupling is outside the exact separable theorem. A negative remaining budget cannot be fixed by silencing these base stations alone.')

s=Slide('A target, an estimate and a confidence bound are different','BACKUP · STATISTICAL INTERPRETATION',source='E4/scene_coverage.csv · E7/strategy_summary.csv')
s.card(44,142,420,250,'E4 scene miscoverage','Target: 10%\nObserved: 13 / 128 = 10.16%\n95% interval: 5.52%–16.74%\n\nCompatible with target; not proof of a ≤10% physical risk.',LIGHT)
s.card(494,142,422,250,'E7: J at Γ = −15 dB','Model risk bound: 8%\nObserved exposure: 8.19%\n95% bootstrap: 6.67%–9.75%\n\n400 complete trajectories; ticks are not independent samples.',PALE)
s.text('Split and resample whole scenes / trajectories. Keep missed and silent receivers.\nZero observed events do not establish zero risk.',54,421,851,19,RED,bold=True)
s.note('The E4 interval is Clopper–Pearson for independent scene failure indicators. There are 128 independent scenes even though the stratified CSV has many more rows. The E7 interval resamples complete episodes because time samples within an episode are dependent. These are different experiments: E4’s scene marginal coverage does not certify E7’s conditional state risks. A measured value just above a budget is not by itself evidence of a solver or theorem failure when the confidence interval includes the target. Conversely, a mean below budget is not a high-confidence certification if its upper confidence bound exceeds that budget. Formal simultaneous or one-sided risk certification needs an explicitly designed statistical procedure.')

s=Slide('Notation and source map','BACKUP · REFERENCE',source='Local paper and pinned result archive; full paths and hashes in sources.json')
items=[('v = √p w','Power-weighted transmit beam; actual power = Pmax p'),('Aᵣ, Cᵣ, ρᵣ','DL path responses, coefficient bound, residual radius'),('Γ / δNTN','INR magnitude limit / active-time exceedance budget'),('L / δTN','Frozen deadline demand / service-failure budget'),('η / c / d','Conditional coverage-risk bound / risk time / active time'),('S / x(S,a) / π(a|S)','Observable history state / occupation mass / action probability')]
y=136
for i,(term,meaning) in enumerate(items):
    s.rect(44,y,872,44,LIGHT if i%2==0 else WHITE,r=4)
    s.text(term,56,y+12,255,16,TEAL,bold=True);s.text(meaning,312,y+12,590,16);y+=48
s.text('Paper: overleaf_ntn_paper/direction.tex\nResults: result/appendix_d_20260921_235448_775259/',55,441,850,14,MUTED)
s.note('This is a reference slide for questions. Sources are local and pinned to the current paper and the September 21 23:54:48 experiment archive. All chart values are replotted from saved CSV or JSON without re-running or modifying the experiments. Mathematical illustrations are explicitly marked schematic. The deck distinguishes current implementation details from the broader Appendix D theorem. Speaker notes cite the relevant local source and explain qualifications that would overload the visible slides. The source manifest records SHA256 hashes so the deck can be reproduced and checked after future changes.')

# ---------------------- renderers ----------------------
def wrap_lines(t,w,size,bold=False):
    font='DV-Bold' if bold else 'DV';out=[]
    for paragraph in t.split('\n'):
        if not paragraph:out.append('');continue
        words=paragraph.split();line=''
        for word in words:
            new=(line+' '+word).strip()
            if line and pdfmetrics.stringWidth(new,font,size)>w:out.append(line);line=word
            else:line=new
        out.append(line)
    return out

def fit_image(path,x,y,w,h):
    im=Image.open(path); iw,ih=im.size;factor=min(w/iw,h/ih)
    ww,hh=iw*factor,ih*factor
    return x+(w-ww)/2,y+(h-hh)/2,ww,hh

def rgb(c):return RGBColor.from_string(c.lstrip('#'))

prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
prs.core_properties.title='Appendix D — Service-Constrained Joint Sensing and Robust Transmission'
prs.core_properties.subject='English explanatory slides, with theory, diagrams, empirical results and speaker notes'
prs.core_properties.author='NTN / TN Coexistence Research'
pdf=canvas.Canvas(str(HERE/'Appendix_D_Explained.pdf'),pagesize=(W,H))
pdf.setTitle(prs.core_properties.title);pdf.setAuthor('NTN / TN Coexistence Research')
layout_warnings=[]
for index,s in enumerate(slides,1):
    ps=prs.slides.add_slide(prs.slide_layouts[6])
    for e in s.e:
        k=e['kind'];x,y=e.get('x',0),e.get('y',0)
        if k=='rect':
            pdf.setFillColor(HexColor(e['color']));pdf.setStrokeColor(HexColor(e['line'] or e['color']))
            pdf.roundRect(x,H-y-e['h'],e['w'],e['h'],e['r'],stroke=bool(e['line']),fill=1)
            sh=ps.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e['r'] else MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(e['w']),Pt(e['h']))
            sh.fill.solid();sh.fill.fore_color.rgb=rgb(e['color']);sh.line.fill.background()
            if e['r']:
                try:sh.adjustments[0]=.08
                except Exception:pass
        elif k=='text':
            lines=wrap_lines(e['t'],e['w'],e['size'],e['bold'])
            height=len(lines)*e['size']*e['leading']+5
            if y+height>H-22 and not s.dark:layout_warnings.append(f'Slide {index}: text reaches footer: {e["t"][:40]}')
            pdf.setFont('DV-Bold' if e['bold'] else 'DV',e['size']);pdf.setFillColor(HexColor(e['color']))
            for j,line in enumerate(lines):
                base=H-y-e['size']*.94-j*e['size']*e['leading']
                if e['align']=='center':pdf.drawCentredString(x+e['w']/2,base,line)
                else:pdf.drawString(x,base,line)
            sh=ps.shapes.add_textbox(Pt(x),Pt(y),Pt(e['w']),Pt(height));tf=sh.text_frame
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
            for j,line in enumerate(lines):
                p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line
                p.font.name='DejaVu Sans';p.font.size=Pt(e['size']);p.font.bold=e['bold'];p.font.color.rgb=rgb(e['color'])
                p.line_spacing=Pt(e['size']*e['leading']);p.space_before=p.space_after=Pt(0)
                if e['align']=='center':
                    from pptx.enum.text import PP_ALIGN
                    p.alignment=PP_ALIGN.CENTER
        elif k=='line':
            pdf.setStrokeColor(HexColor(e['color']));pdf.setLineWidth(e['width']);pdf.setDash(3,3) if e['dash'] else pdf.setDash()
            pdf.line(x,H-y,e['x2'],H-e['y2']);pdf.setDash()
            sh=ps.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x),Pt(y),Pt(e['x2']),Pt(e['y2']))
            sh.line.color.rgb=rgb(e['color']);sh.line.width=Pt(e['width'])
        elif k=='triangle':
            pts=e['points'];pdf.setFillColor(HexColor(e['color']));p=pdf.beginPath();p.moveTo(pts[0][0],H-pts[0][1])
            for a,b in pts[1:]:p.lineTo(a,H-b)
            p.close();pdf.drawPath(p,stroke=0,fill=1)
            # Native freeform arrowhead, not a raster screenshot.
            builder=ps.shapes.build_freeform(Pt(pts[0][0]),Pt(pts[0][1]),scale=1)
            builder.add_line_segments([(Pt(a),Pt(b)) for a,b in pts[1:]],close=True)
            sh=builder.convert_to_shape();sh.fill.solid();sh.fill.fore_color.rgb=rgb(e['color']);sh.line.fill.background()
        elif k=='image':
            xx,yy,ww,hh=fit_image(e['path'],x,y,e['w'],e['h'])
            pdf.drawImage(e['path'],xx,H-yy-hh,ww,hh,mask='auto')
            ps.shapes.add_picture(e['path'],Pt(xx),Pt(yy),Pt(ww),Pt(hh))
    # Footer and notes on every slide.
    color='#BCD0DD' if s.dark else MUTED
    footer=s.source
    if pdfmetrics.stringWidth(footer,'DV',9)>840:
        footer=footer[:145]+'…'
    pdf.setFont('DV',9);pdf.setFillColor(HexColor(color));pdf.drawString(44,15,footer);pdf.drawRightString(916,15,f'{index:02d} / {len(slides)}')
    sh=ps.shapes.add_textbox(Pt(44),Pt(516),Pt(835),Pt(17));p=sh.text_frame.paragraphs[0];p.text=footer
    sh.text_frame.margin_left=sh.text_frame.margin_top=0;p.font.name='DejaVu Sans';p.font.size=Pt(9);p.font.color.rgb=rgb(color)
    sh=ps.shapes.add_textbox(Pt(885),Pt(516),Pt(65),Pt(17));p=sh.text_frame.paragraphs[0];p.text=f'{index:02d} / {len(slides)}'
    sh.text_frame.margin_left=sh.text_frame.margin_top=0;p.font.name='DejaVu Sans';p.font.size=Pt(9);p.font.color.rgb=rgb(color)
    ps.notes_slide.notes_text_frame.text=f'Slide {index}: {s.title}\n\n{s.notes}\n\nSource: {s.source}'
    pdf.showPage()
pdf.save();prs.save(HERE/'Appendix_D_Explained.pptx')

md=['# Appendix D — English speaker notes','',
    'Main presentation: slides 1–25, approximately 25–30 minutes. Backup: slides 26–30.',
    'The editable PowerPoint uses native text and diagram shapes; equations and scientific plots are high-resolution images.',
    'Paper source: `overleaf_ntn_paper/direction.tex`. Pinned run: `result/appendix_d_20260921_235448_775259`.', '']
for i,s in enumerate(slides,1):md += [f'## {i:02d}. {s.title}','',s.notes,'',f'Source: {s.source}','']
(HERE/'Speaker_Notes.md').write_text('\n'.join(md)+'\n')
styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='TalkTitle',fontName='DV-Bold',fontSize=18,leading=24,textColor=HexColor(NAVY),spaceAfter=20))
styles.add(ParagraphStyle(name='TalkBody',fontName='DV',fontSize=11,leading=17,spaceAfter=14))
styles.add(ParagraphStyle(name='TalkSource',fontName='DV',fontSize=9,leading=13,textColor=HexColor(MUTED)))
story=[]
for i,s in enumerate(slides,1):
    if i>1:story.append(PageBreak())
    story.append(Paragraph(escape(f'{i:02d}. {s.title}'),styles['TalkTitle']))
    story.append(Paragraph(escape(s.notes),styles['TalkBody']))
    story.append(Spacer(1,10));story.append(Paragraph(escape('Source: '+s.source),styles['TalkSource']))
SimpleDocTemplate(str(HERE/'Speaker_Notes.pdf'),pagesize=(595,842),leftMargin=48,rightMargin=48,topMargin=56,bottomMargin=50).build(story)
files=[PAPER,RUN/'manifest.json',RUN/'E7/strategy_summary.csv',RUN/'E6/E6_validation.csv',RUN/'E4/scene_coverage.csv',RUN/'fresh_rt_spatial/gamma_summary.csv',RUN/'E8/random_arrival/summary.json',RUN/'E8/repeated_observation/summary.json']
(HERE/'sources.json').write_text(json.dumps({'paper':str(PAPER.relative_to(ROOT)),'run':str(RUN.relative_to(ROOT)),
    'slides':len(slides),'main_slides':25,'backup_slides':5,'sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in files},
    'diagram_scope':'timing, architecture and uncertainty diagrams are conceptual; charts use the pinned result files',
    'layout_warnings':layout_warnings},indent=2)+'\n')
(HERE/'slide_content.json').write_text(json.dumps([dict(number=i,title=s.title,section=s.section,notes=s.notes,source=s.source) for i,s in enumerate(slides,1)],indent=2)+'\n')
print(f'Built {len(slides)} slides: PDF, editable PPTX, and speaker notes.')
print('Layout warnings:',layout_warnings)
